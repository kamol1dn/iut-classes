#!/usr/bin/env python3
"""Dump the text of a lecture deck so it can be read, searched or summarised.

Handles the formats the courses actually use:

    .pptx   python-pptx — slide text, tables, and speaker notes
    .pdf    pypdf — page text

Old binary ``.pps`` / ``.ppt`` are NOT supported by python-pptx. Convert first:
``soffice --headless --convert-to pptx <file>``.

    python scripts/extract_slides.py "path/to/deck.pptx"
    python scripts/extract_slides.py "path/to/deck.pdf" -o out.txt
    python scripts/extract_slides.py deck.pptx --slides 5-12
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")


def parse_range(spec: str | None, total: int) -> range:
    if not spec:
        return range(1, total + 1)
    if "-" in spec:
        lo, hi = spec.split("-", 1)
        return range(int(lo), min(int(hi), total) + 1)
    n = int(spec)
    return range(n, n + 1)


def from_pptx(path: Path, spec: str | None) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = list(prs.slides)
    out = []
    for i in parse_range(spec, len(slides)):
        slide = slides[i - 1]
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        try:
            if slide.has_notes_slide:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    parts.append(f"[NOTES] {note}")
        except Exception:
            pass
        body = "\n".join(parts).strip()
        out.append(f"--- slide {i} ---\n{body}" if body else f"--- slide {i} --- (no text)")
    return out


def from_pdf(path: Path, spec: str | None) -> list[str]:
    import pypdf

    reader = pypdf.PdfReader(str(path))
    pages = reader.pages
    out = []
    for i in parse_range(spec, len(pages)):
        text = (pages[i - 1].extract_text() or "").strip()
        out.append(f"--- page {i} ---\n{text}" if text else f"--- page {i} --- (no text)")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("file", type=Path)
    ap.add_argument("-o", "--out", type=Path)
    ap.add_argument("--slides", "--pages", dest="slides",
                    help="e.g. 5 or 5-12 (default: all)")
    args = ap.parse_args()

    if not args.file.exists():
        print(f"not found: {args.file}", file=sys.stderr)
        return 1

    ext = args.file.suffix.lower()
    if ext == ".pptx":
        blocks = from_pptx(args.file, args.slides)
    elif ext == ".pdf":
        blocks = from_pdf(args.file, args.slides)
    elif ext in (".pps", ".ppt"):
        print(f"{ext} is the old binary format — convert it first:\n"
              f'  soffice --headless --convert-to pptx "{args.file}"', file=sys.stderr)
        return 2
    else:
        print(f"unsupported: {ext}", file=sys.stderr)
        return 2

    text = f"# {args.file.name}\n\n" + "\n\n".join(blocks)
    if args.out:
        args.out.write_text(text, encoding="utf-8")
        print(f"{len(blocks)} block(s) -> {args.out}")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
